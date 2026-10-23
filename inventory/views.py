import io
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.db.models import Q
from django.utils import timezone
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger

from .models import Product, Category, CategoryType, ProductBatch
from .forms import ProductForm, StockMovementForm, AddBatchForm, RemoveStockForm, CategoryForm, CategoryTypeForm
from activity.models import ActivityLog


# ── PRODUCTS ──────────────────────────────────────────────────────────────────

@login_required
def product_list(request):
    query       = request.GET.get('q', '')
    category_id = request.GET.get('category', '')
    status_filter = request.GET.get('status', '')

    products = Product.objects.filter(is_active=True).select_related('category__category_type')

    if query:
        products = products.filter(Q(name__icontains=query) | Q(category__name__icontains=query))
    if category_id:
        products = products.filter(category_id=category_id)
    if status_filter:
        products = products.filter(status=status_filter)

    total_count = products.count()
    paginator = Paginator(products, 10)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    categories = Category.objects.all()
    context = {
        'products': page_obj, 'page_obj': page_obj,
        'paginator': paginator, 'total_count': total_count,
        'categories': categories,
        'query': query, 'selected_category': category_id, 'selected_status': status_filter,
    }
    return render(request, 'inventory/product_list.html', context)


@login_required
def live_search(request):
    """AJAX endpoint — handles search text + category + status filters, no page reload."""
    from django.utils import timezone as tz
    query         = request.GET.get('q', '').strip()
    category_id   = request.GET.get('category', '').strip()
    status_filter = request.GET.get('status', '').strip()
    today         = tz.now().date()

    products = Product.objects.filter(is_active=True).select_related('category')

    if query:
        products = products.filter(
            Q(name__icontains=query) | Q(category__name__icontains=query)
        )
    if category_id:
        products = products.filter(category_id=category_id)
    if status_filter:
        products = products.filter(status=status_filter)

    data = {
        "products": [
            {
                "pk": p.pk,
                "name": p.name,
                "category": p.category.name if p.category else "-",
                "quantity": p.quantity,
                "status": p.status,
                "expiry_date": str(p.expiry_date) if p.expiry_date else "-",
                "is_expired": bool(p.expiry_date and p.expiry_date < today),
                "is_expiring_soon": bool(
                    p.expiry_date and
                    today <= p.expiry_date and
                    (p.expiry_date - today).days <= 7
                ),
            }
            for p in products[:100]
        ]
    }
    return JsonResponse(data)


@login_required
def live_search_inactive_products(request):
    """AJAX endpoint for inactive products search."""
    from django.utils import timezone as tz
    query = request.GET.get('q', '').strip()
    today = tz.now().date()

    products = Product.objects.filter(is_active=False).select_related('category')
    if query:
        products = products.filter(
            Q(name__icontains=query) | Q(category__name__icontains=query)
        )

    data = {
        'products': [
            {
                'pk': p.pk,
                'name': p.name,
                'category': p.category.name if p.category else '-',
                'quantity': p.quantity,
                'status': p.status,
                'expiry_date': str(p.expiry_date) if p.expiry_date else '-',
                'is_expired': bool(p.expiry_date and p.expiry_date < today),
                'deactivated_at': p.deactivated_at.strftime('%b %d, %Y %I:%M %p') if p.deactivated_at else '—',
            }
            for p in products[:100]
        ]
    }
    return JsonResponse(data)


@login_required
def live_search_inactive_categories(request):
    """AJAX endpoint for inactive categories search."""
    query = request.GET.get('q', '').strip()

    categories = Category.objects.filter(is_active=False).select_related('category_type')
    if query:
        categories = categories.filter(name__icontains=query)

    data = {
        'categories': [
            {
                'pk': c.pk,
                'name': c.name,
                'color': c.color,
                'icon': c.icon,
                'category_type': c.category_type.name if c.category_type else None,
                'product_count': c.products.count(),
                'description': c.description or '',
                'deactivated_at': c.deactivated_at.strftime('%b %d, %Y %I:%M %p') if c.deactivated_at else '—',
            }
            for c in categories[:100]
        ]
    }
    return JsonResponse(data)


@login_required
def live_search_categories(request):
    """AJAX endpoint for active categories search."""
    query = request.GET.get('q', '').strip()

    categories = Category.objects.filter(is_active=True).select_related('category_type')
    if query:
        categories = categories.filter(name__icontains=query)

    data = {
        'categories': [
            {
                'pk': c.pk,
                'name': c.name,
                'color': c.color,
                'icon': c.icon,
                'category_type': c.category_type.name if c.category_type else None,
                'product_count': c.products.count(),
                'description': c.description or '',
            }
            for c in categories[:100]
        ]
    }
    return JsonResponse(data)


@login_required
def inactive_products(request):
    query = request.GET.get('q', '')
    products = Product.objects.filter(is_active=False).select_related('category__category_type')
    if query:
        products = products.filter(Q(name__icontains=query) | Q(category__name__icontains=query))
    total_count = products.count()
    paginator = Paginator(products, 10)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)
    return render(request, 'inventory/inactive_products.html', {
        'products': page_obj, 'page_obj': page_obj,
        'paginator': paginator, 'total_count': total_count, 'query': query,
    })


@login_required
def clear_inactive_products(request):
    """Fix 5: Permanently delete all inactive products; preserve activity logs."""
    if request.method == 'POST':
        inactive_qs = Product.objects.filter(is_active=False)
        count = inactive_qs.count()
        if count == 0:
            messages.info(request, 'There are no inactive products to clear.')
            return redirect('inactive_products')
        # Detach activity logs from products before deletion so logs are preserved
        from activity.models import ActivityLog
        ActivityLog.objects.filter(product__in=inactive_qs).update(product=None)
        inactive_qs.delete()
        # Record the cleanup action (no product reference — it's a system-level action)
        ActivityLog.objects.create(
            product=None,
            product_name='[System]',
            user=request.user,
            action='DELETE',
            quantity=count,
            notes=f'Cleared {count} inactive product(s) permanently. Activity logs preserved.'
        )
        messages.success(request, f'{count} inactive product(s) have been permanently removed. Activity logs are preserved.')
        return redirect('inactive_products')
    # GET: show confirmation page
    count = Product.objects.filter(is_active=False).count()
    return render(request, 'inventory/clear_inactive_confirm.html', {'count': count})


@login_required
def product_deactivate(request, pk):
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        product.deactivate()
        ActivityLog.objects.create(
            product=product, user=request.user,
            action='DELETE', quantity=0, notes='Product deactivated (soft delete)'
        )
        messages.success(request, f'"{product.name}" has been deactivated.')
        return redirect('product_list')
    return render(request, 'inventory/product_confirm_deactivate.html', {'product': product})


@login_required
def product_restore_confirm(request, pk):
    product = get_object_or_404(Product, pk=pk, is_active=False)
    category_inactive = product.category and not product.category.is_active
    return render(request, 'inventory/product_confirm_restore.html', {
        'product': product,
        'category_inactive': category_inactive,
    })


@login_required
def product_restore(request, pk):
    product = get_object_or_404(Product, pk=pk, is_active=False)
    # Dependency check: block restore if the product's category is inactive
    if product.category and not product.category.is_active:
        messages.error(
            request,
            f'Cannot restore "{product.name}" because its category '
            f'"{product.category.name}" is still inactive. '
            f'Please restore the category first.'
        )
        return redirect('inactive_products')
    # Fix 3: Do NOT recreate batches — only flip the active flag
    product.is_active = True
    product.deactivated_at = None
    product.save(update_fields=['is_active', 'deactivated_at', 'last_updated'])
    ActivityLog.objects.create(
        product=product, user=request.user,
        action='UPDATE', quantity=0, notes='Product restored from inactive'
    )
    messages.success(request, f'"{product.name}" has been restored.')
    return redirect('inactive_products')


@login_required
def product_add(request):
    if request.method == 'POST':
        form = ProductForm(request.POST)
        if form.is_valid():
            name = form.cleaned_data['name'].strip()
            # Fix 4: Duplicate product name check (case-insensitive)
            existing = Product.objects.filter(name__iexact=name).first()
            if existing:
                if existing.is_active:
                    messages.error(request, f'A product named "{existing.name}" already exists. Please update the existing product instead.')
                else:
                    messages.error(request, f'A product named "{existing.name}" already exists but is inactive. Please restore it from Inactive Products.')
                return render(request, 'inventory/product_form.html', {'form': form, 'action': 'Add'})
            product = form.save()
            # Create the initial batch from the product's quantity and expiry_date
            ProductBatch.objects.create(
                product=product,
                quantity=product.quantity,
                expiry_date=product.expiry_date,
                added_by=request.user,
            )
            ActivityLog.objects.create(
                product=product, user=request.user,
                action='ADD', quantity=product.quantity,
                notes=f'Product created – initial batch (Exp: {product.expiry_date})'
            )
            messages.success(request, f'Product "{product.name}" added successfully!')
            return redirect('product_list')
    else:
        form = ProductForm()
    return render(request, 'inventory/product_form.html', {'form': form, 'action': 'Add'})


@login_required
def product_edit(request, pk):
    product = get_object_or_404(Product, pk=pk)
    old_quantity = product.quantity
    if request.method == 'POST':
        form = ProductForm(request.POST, instance=product)
        if form.is_valid():
            name = form.cleaned_data['name'].strip()
            # Fix 4: Check for duplicate name (exclude current product)
            existing = Product.objects.filter(name__iexact=name).exclude(pk=pk).first()
            if existing:
                messages.error(request, f'Another product named "{existing.name}" already exists.')
                return render(request, 'inventory/product_form.html', {'form': form, 'action': 'Edit', 'product': product})
            updated = form.save()
            new_quantity = updated.quantity
            diff = new_quantity - old_quantity

            # Batch sync: adjust batch quantities to match the corrected product total
            if diff != 0:
                batches = list(updated.batches.order_by('date_added', 'id'))
                if diff > 0:
                    # Quantity increased: add difference to most recent batch (or create one)
                    if batches:
                        latest_batch = batches[-1]
                        latest_batch.quantity += diff
                        latest_batch.save()
                    else:
                        from .models import ProductBatch
                        ProductBatch.objects.create(
                            product=updated,
                            quantity=new_quantity,
                            expiry_date=updated.expiry_date,
                            added_by=request.user,
                        )
                else:
                    # Quantity decreased: deduct FIFO from oldest batch(es)
                    remaining = abs(diff)
                    for batch in batches:
                        if remaining <= 0:
                            break
                        deduct = min(remaining, batch.quantity)
                        batch.quantity -= deduct
                        remaining -= deduct
                        if batch.quantity == 0:
                            batch.delete()
                        else:
                            batch.save()

                ActivityLog.objects.create(
                    product=updated, user=request.user,
                    action='ADD' if diff > 0 else 'REMOVE',
                    quantity=abs(diff),
                    notes=f'Product edited – quantity {"increased" if diff > 0 else "decreased"} by {abs(diff)} (batch synced)'
                )

            messages.success(request, f'Product "{updated.name}" updated!')
            return redirect('product_list')
    else:
        form = ProductForm(instance=product)
    return render(request, 'inventory/product_form.html', {'form': form, 'action': 'Edit', 'product': product})


@login_required
def product_detail(request, pk):
    product = get_object_or_404(Product, pk=pk)
    logs    = ActivityLog.objects.filter(product=product).order_by('-date')[:20]
    batches = product.batches.all().order_by('date_added', 'id')
    return render(request, 'inventory/product_detail.html', {
        'product': product,
        'logs': logs,
        'batches': batches,
    })


@login_required
def stock_movement(request, pk):
    """Legacy redirect – routes to add_stock page."""
    return redirect('add_stock', pk=pk)


@login_required
def add_stock(request, pk):
    """Create a new batch for the product (FIFO-aware stock addition)."""
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        form = AddBatchForm(request.POST)
        if form.is_valid():
            qty         = form.cleaned_data['quantity']
            expiry_date = form.cleaned_data['expiry_date']
            notes       = form.cleaned_data.get('notes', '')

            # Create new batch – do NOT touch existing batches
            batch = ProductBatch.objects.create(
                product=product,
                quantity=qty,
                expiry_date=expiry_date,
                added_by=request.user,
            )
            # Sync total quantity back to Product
            product.sync_quantity()

            ActivityLog.objects.create(
                product=product, user=request.user,
                action='ADD', quantity=qty,
                notes=f'New batch #{batch.pk} added | Exp: {expiry_date}' + (f' | {notes}' if notes else '')
            )
            messages.success(request, f'Added {qty} units to "{product.name}" as a new batch (Exp: {expiry_date}).')
            return redirect('product_detail', pk=product.pk)
    else:
        form = AddBatchForm()
    return render(request, 'inventory/add_stock.html', {'form': form, 'product': product})


@login_required
def remove_stock(request, pk):
    """Remove stock using FIFO – deduct from the oldest batch first."""
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        form = RemoveStockForm(request.POST)
        if form.is_valid():
            qty_to_remove = form.cleaned_data['quantity']
            notes         = form.cleaned_data.get('notes', '')
            total         = product.total_quantity

            if qty_to_remove > total:
                messages.error(request, f'Cannot remove {qty_to_remove}. Only {total} units across all batches.')
                return render(request, 'inventory/remove_stock.html', {'form': form, 'product': product})

            # FIFO: deduct from oldest batches first
            remaining      = qty_to_remove
            affected_info  = []
            batches        = product.batches.order_by('date_added', 'id')
            for batch in batches:
                if remaining <= 0:
                    break
                deduct = min(remaining, batch.quantity)
                batch.quantity -= deduct
                remaining -= deduct
                affected_info.append(f'Batch #{batch.pk} (Exp: {batch.expiry_date}) -{deduct}')
                if batch.quantity == 0:
                    batch.delete()
                else:
                    batch.save()

            product.sync_quantity()
            # Fix 1: Dynamically update product expiry after batch removal
            product.sync_expiry()

            ActivityLog.objects.create(
                product=product, user=request.user,
                action='REMOVE', quantity=qty_to_remove,
                notes=f'FIFO removal | {"; ".join(affected_info)}' + (f' | {notes}' if notes else '')
            )
            messages.success(request, f'Removed {qty_to_remove} units from "{product.name}" (FIFO).')
            return redirect('product_detail', pk=product.pk)
    else:
        form = RemoveStockForm()
    batches = product.batches.order_by('date_added', 'id')
    return render(request, 'inventory/remove_stock.html', {'form': form, 'product': product, 'batches': batches})



@login_required
def batch_edit(request, product_pk, batch_pk):
    """Edit a specific batch's quantity and expiry date."""
    from .models import ProductBatch
    from .forms import BatchEditForm
    product = get_object_or_404(Product, pk=product_pk)
    batch   = get_object_or_404(ProductBatch, pk=batch_pk, product=product)

    if request.method == 'POST':
        form = BatchEditForm(request.POST)
        if form.is_valid():
            old_quantity    = batch.quantity
            new_quantity    = form.cleaned_data['quantity']
            new_expiry      = form.cleaned_data['expiry_date']
            old_expiry      = batch.expiry_date

            batch.quantity    = new_quantity
            batch.expiry_date = new_expiry
            batch.save()

            # Sync product total quantity and expiry
            product.sync_quantity()
            product.sync_expiry()

            # Build a meaningful activity note
            changes = []
            if old_quantity != new_quantity:
                changes.append(f'qty {old_quantity} → {new_quantity}')
            if old_expiry != new_expiry:
                changes.append(f'expiry {old_expiry} → {new_expiry}')
            note = f'Batch #{batch.pk} edited: {", ".join(changes)}' if changes else f'Batch #{batch.pk} saved (no changes)'

            ActivityLog.objects.create(
                product=product,
                user=request.user,
                action='UPDATE',
                quantity=new_quantity,
                notes=note,
            )
            messages.success(request, f'Batch #{batch.pk} updated successfully. Product total is now {product.quantity} pcs.')
            return redirect('product_detail', pk=product.pk)
    else:
        form = BatchEditForm(initial={
            'quantity':    batch.quantity,
            'expiry_date': batch.expiry_date,
        })

    # All batches for context panel
    batches = product.batches.order_by('date_added', 'id')
    return render(request, 'inventory/batch_edit.html', {
        'form':    form,
        'product': product,
        'batch':   batch,
        'batches': batches,
    })


# ── CATEGORIES ────────────────────────────────────────────────────────────────

@login_required
def category_list(request):
    selected_type = request.GET.get('type', '')
    query         = request.GET.get('q', '')
    categories    = Category.objects.filter(is_active=True).select_related('category_type').all()
    if selected_type:
        categories = categories.filter(category_type__name=selected_type)
    if query:
        categories = categories.filter(name__icontains=query)
    cat_types = CategoryType.objects.all()
    inactive_count = Category.objects.filter(is_active=False).count()
    total_count = categories.count()
    paginator = Paginator(categories, 12)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)
    context = {
        'categories':     page_obj,
        'page_obj':       page_obj,
        'paginator':      paginator,
        'total_count':    total_count,
        'cat_types':      cat_types,
        'selected_type':  selected_type,
        'inactive_count': inactive_count,
        'query':          query,
    }
    return render(request, 'inventory/category_list.html', context)


@login_required
def inactive_categories(request):
    query = request.GET.get('q', '')
    categories = Category.objects.filter(is_active=False).select_related('category_type')
    if query:
        categories = categories.filter(name__icontains=query)
    total_count = categories.count()
    paginator = Paginator(categories, 12)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)
    return render(request, 'inventory/inactive_categories.html', {
        'categories': page_obj, 'page_obj': page_obj,
        'paginator': paginator, 'total_count': total_count, 'query': query,
    })


@login_required
def clear_inactive_categories(request):
    """Permanently delete all inactive categories and their related inactive products."""
    if request.method == 'POST':
        inactive_qs = Category.objects.filter(is_active=False)
        count = inactive_qs.count()
        if count == 0:
            messages.info(request, 'There are no inactive categories to clear.')
            return redirect('inactive_categories')
        # Also delete related inactive products (orphan prevention)
        from .models import Product as _Product
        product_qs = _Product.objects.filter(category__in=inactive_qs, is_active=False)
        product_count = product_qs.count()
        # Detach activity logs from those products so logs are preserved
        from activity.models import ActivityLog
        ActivityLog.objects.filter(product__in=product_qs).update(product=None)
        product_qs.delete()
        inactive_qs.delete()
        ActivityLog.objects.create(
            product=None,
            product_name='[System]',
            user=request.user,
            action='DELETE',
            quantity=count,
            notes=(
                f'Cleared {count} inactive category(ies) and {product_count} '
                f'related inactive product(s) permanently. Activity logs preserved.'
            )
        )
        messages.success(
            request,
            f'{count} inactive category(ies) and {product_count} related inactive '
            f'product(s) have been permanently removed. Activity logs are preserved.'
        )
        return redirect('inactive_categories')
    # GET: show confirmation page
    count = Category.objects.filter(is_active=False).count()
    from .models import Product as _Product
    product_count = _Product.objects.filter(
        category__in=Category.objects.filter(is_active=False), is_active=False
    ).count()
    return render(request, 'inventory/clear_inactive_categories_confirm.html', {
        'count': count, 'product_count': product_count,
    })


@login_required
def category_restore_confirm(request, pk):
    category = get_object_or_404(Category, pk=pk, is_active=False)
    product_count = category.products.filter(is_active=False).count()
    return render(request, 'inventory/category_confirm_restore.html', {
        'category': category, 'product_count': product_count,
    })


@login_required
def category_restore(request, pk):
    category = get_object_or_404(Category, pk=pk, is_active=False)
    if request.method == 'POST':
        product_count = category.products.filter(is_active=False).count()
        category.restore()
        ActivityLog.objects.create(
            product=None, product_name='[System]', user=request.user,
            action='UPDATE', quantity=0,
            notes=f'Category "{category.name}" restored with {product_count} product(s).'
        )
        messages.success(request, f'Category "{category.name}" and {product_count} product(s) have been restored.')
        return redirect('inactive_categories')
    return redirect('category_restore_confirm', pk=pk)


@login_required
def category_add(request):
    if request.method == 'POST':
        form = CategoryForm(request.POST)
        if form.is_valid():
            name = form.cleaned_data['name'].strip()
            existing = Category.objects.filter(name__iexact=name).first()
            if existing:
                if existing.is_active:
                    messages.error(request, f'A category named "{existing.name}" already exists. Please update it instead.')
                else:
                    messages.error(request, f'A category named "{existing.name}" is inactive. Please restore it from Inactive Categories.')
            else:
                form.save()
                messages.success(request, f'Category "{name}" added!')
                return redirect('category_list')
    else:
        form = CategoryForm()
    cat_types = CategoryType.objects.all()
    return render(request, 'inventory/category_form.html', {'form': form, 'action': 'Add', 'cat_types': cat_types})


@login_required
def category_edit(request, pk):
    category = get_object_or_404(Category, pk=pk)
    if request.method == 'POST':
        form = CategoryForm(request.POST, instance=category)
        if form.is_valid():
            name = form.cleaned_data['name'].strip()
            existing = Category.objects.filter(name__iexact=name).exclude(pk=pk).first()
            if existing:
                messages.error(request, f'Another category named "{existing.name}" already exists.')
            else:
                form.save()
                messages.success(request, f'Category "{name}" updated!')
                return redirect('category_list')
    else:
        form = CategoryForm(instance=category)
    cat_types = CategoryType.objects.all()
    return render(request, 'inventory/category_form.html', {
        'form': form, 'action': 'Edit', 'category': category, 'cat_types': cat_types
    })


@login_required
def category_delete(request, pk):
    category = get_object_or_404(Category, pk=pk, is_active=True)
    if request.method == 'POST':
        product_count = category.products.filter(is_active=True).count()
        category.deactivate()
        ActivityLog.objects.create(
            product=None, product_name='[System]', user=request.user,
            action='DELETE', quantity=0,
            notes=f'Category "{category.name}" deactivated with {product_count} product(s).'
        )
        messages.success(request, f'Category "{category.name}" and {product_count} product(s) moved to inactive.')
        return redirect('category_list')
    product_count = category.products.filter(is_active=True).count()
    return render(request, 'inventory/category_confirm_delete.html', {
        'category': category, 'product_count': product_count,
    })


# ── CATEGORY TYPES ────────────────────────────────────────────────────────────

@login_required
def category_type_list(request):
    types = CategoryType.objects.all()
    return render(request, 'inventory/category_type_list.html', {'types': types})


@login_required
def category_type_add(request):
    if request.method == 'POST':
        form = CategoryTypeForm(request.POST)
        if form.is_valid():
            ct = form.save()
            messages.success(request, f'Category type "{ct.name}" created!')
            # If AJAX, return JSON for live UI update
            if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                return JsonResponse({'id': ct.pk, 'name': ct.name, 'color': ct.color, 'icon': ct.icon})
            return redirect('category_type_list')
        if request.headers.get('x-requested-with') == 'XMLHttpRequest':
            return JsonResponse({'error': form.errors}, status=400)
    else:
        form = CategoryTypeForm()
    return render(request, 'inventory/category_type_form.html', {'form': form, 'action': 'Add'})


@login_required
def category_type_edit(request, pk):
    ct = get_object_or_404(CategoryType, pk=pk)
    if request.method == 'POST':
        form = CategoryTypeForm(request.POST, instance=ct)
        if form.is_valid():
            form.save()
            messages.success(request, f'Category type "{ct.name}" updated!')
            return redirect('category_type_list')
    else:
        form = CategoryTypeForm(instance=ct)
    return render(request, 'inventory/category_type_form.html', {'form': form, 'action': 'Edit', 'ct': ct})


@login_required
def category_type_delete(request, pk):
    ct = get_object_or_404(CategoryType, pk=pk)
    if request.method == 'POST':
        ct.delete()
        messages.success(request, f'Category type "{ct.name}" deleted.')
        return redirect('category_type_list')
    return render(request, 'inventory/category_type_confirm_delete.html', {'ct': ct})


# ── EXPORTS ───────────────────────────────────────────────────────────────────

@login_required
def export_inventory_excel(request):
    products = Product.objects.filter(is_active=True).select_related('category__category_type')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory Report"

    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    headers = ['#', 'Product Name', 'Category', 'Type', 'Quantity', 'Status', 'Expiry Date']

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.row_dimensions[1].height = 22
    for i, w in enumerate([5, 28, 18, 14, 12, 15, 16], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w

    status_colors = {'In Stock': 'DCFCE7', 'Low Stock': 'FEF9C3', 'Out of Stock': 'FEE2E2'}
    for row_num, p in enumerate(products, 2):
        row = [row_num-1, p.name, p.category.name,
               p.category.type_name, p.quantity, p.status, str(p.expiry_date)]
        for col, val in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.alignment = Alignment(horizontal='center' if col != 2 else 'left', vertical='center')
            if col == 6:
                cell.fill = PatternFill("solid", fgColor=status_colors.get(p.status, 'FFFFFF'))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="inventory_report.xlsx"'
    wb.save(response)
    return response


@login_required
def export_inventory_pdf(request):
    products = Product.objects.filter(is_active=True).select_related('category__category_type')
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.75*inch, bottomMargin=0.75*inch)
    styles = getSampleStyleSheet()
    story  = []

    title_s = ParagraphStyle('T', parent=styles['Title'], fontSize=18,
                              textColor=colors.HexColor('#1E40AF'), spaceAfter=6)
    sub_s   = ParagraphStyle('S', parent=styles['Normal'], fontSize=10,
                              textColor=colors.grey, spaceAfter=20)
    story.append(Paragraph("Inventory Report", title_s))
    story.append(Paragraph(f"Generated: {timezone.now().strftime('%B %d, %Y %H:%M')}", sub_s))

    data = [['#', 'Product Name', 'Category', 'Type', 'Qty', 'Status', 'Expiry']]
    sc   = {'In Stock': colors.HexColor('#DCFCE7'), 'Low Stock': colors.HexColor('#FEF9C3'),
            'Out of Stock': colors.HexColor('#FEE2E2')}
    for i, p in enumerate(products, 1):
        data.append([str(i), p.name, p.category.name, p.category.type_name,
                     str(p.quantity), p.status, str(p.expiry_date)])

    t = Table(data, colWidths=[0.35*inch, 1.9*inch, 1.3*inch, 1.1*inch, 0.5*inch, 1.0*inch, 1.0*inch])
    ts = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E40AF')),
        ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
        ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE',   (0,0), (-1,-1), 9),
        ('ALIGN',      (0,0), (-1,-1), 'CENTER'),
        ('ALIGN',      (1,1), (1,-1), 'LEFT'),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID',       (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING',    (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]
    for i, p in enumerate(products, 1):
        if p.status in sc:
            ts.append(('BACKGROUND', (5, i), (5, i), sc[p.status]))
    t.setStyle(TableStyle(ts))
    story.append(t)

    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="inventory_report.pdf"'
    return response


@login_required
def export_expiration_excel(request):
    from datetime import timedelta
    today    = timezone.now().date()
    seven    = today + timedelta(days=7)
    expired  = ProductBatch.objects.filter(expiry_date__lt=today, product__is_active=True).select_related('product__category').order_by('expiry_date')
    expiring = ProductBatch.objects.filter(expiry_date__range=[today, seven], product__is_active=True).select_related('product__category').order_by('expiry_date')

    wb = openpyxl.Workbook()
    ws1 = wb.active; ws1.title = "Expired Batches"
    ws2 = wb.create_sheet("Expiring Soon Batches")
    headers = ['#', 'Product', 'Category', 'Type', 'Batch #', 'Quantity', 'Expiry Date']

    for ws, qs, hcolor in [(ws1, expired, "991B1B"), (ws2, expiring, "92400E")]:
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=hcolor)
        for rn, b in enumerate(qs, 2):
            for col, val in enumerate([rn-1, b.product.name, b.product.category.name,
                                        b.product.category.type_name, f'#{b.pk}', b.quantity, str(b.expiry_date)], 1):
                ws.cell(row=rn, column=col, value=val)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="expiration_report.xlsx"'
    wb.save(response)
    return response


@login_required
def export_expiration_pdf(request):
    from datetime import timedelta
    today    = timezone.now().date()
    seven    = today + timedelta(days=7)
    expired  = list(ProductBatch.objects.filter(expiry_date__lt=today, product__is_active=True).select_related('product__category').order_by('expiry_date'))
    expiring = list(ProductBatch.objects.filter(expiry_date__range=[today, seven], product__is_active=True).select_related('product__category').order_by('expiry_date'))

    buffer = io.BytesIO()
    doc    = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.75*inch, bottomMargin=0.75*inch)
    styles = getSampleStyleSheet()
    story  = []

    title_s = ParagraphStyle('T', parent=styles['Title'], fontSize=18,
                              textColor=colors.HexColor('#991B1B'), spaceAfter=4)
    sub_s   = ParagraphStyle('S', parent=styles['Normal'], fontSize=10, textColor=colors.grey, spaceAfter=16)
    h2_s    = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=13,
                              textColor=colors.HexColor('#1E293B'), spaceAfter=8, spaceBefore=16)
    story.append(Paragraph("Expiration Report (Batch Level)", title_s))
    story.append(Paragraph(f"Generated: {timezone.now().strftime('%B %d, %Y %H:%M')}", sub_s))

    def make_table(qs, hcolor):
        data = [['#', 'Product', 'Category', 'Batch #', 'Qty', 'Expiry']]
        for i, b in enumerate(qs, 1):
            data.append([str(i), b.product.name, b.product.category.name, f'#{b.pk}',
                         str(b.quantity), str(b.expiry_date)])
        t = Table(data, colWidths=[0.35*inch, 2.0*inch, 1.3*inch, 0.7*inch, 0.5*inch, 1.0*inch])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor(hcolor)),
            ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
            ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE',   (0,0), (-1,-1), 9),
            ('ALIGN',      (0,0), (-1,-1), 'CENTER'),
            ('ALIGN',      (1,1), (1,-1), 'LEFT'),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID',       (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING',    (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ]))
        return t

    story.append(Paragraph(f"Expired Batches ({len(expired)})", h2_s))
    story.append(make_table(expired, '#991B1B'))
    story.append(Spacer(1, 0.2*inch))
    story.append(Paragraph(f"Expiring Soon ({len(expiring)})", h2_s))
    story.append(make_table(expiring, '#92400E'))

    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="expiration_report.pdf"'
    return response
