from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.db.models import Q
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from .models import LoginActivity
from accounts.views import manager_required


@manager_required
def login_activity_list(request):
    logs = LoginActivity.objects.select_related('user').all()

    query = request.GET.get('q', '')
    user_filter = request.GET.get('user', '')

    if query:
        logs = logs.filter(Q(user__username__icontains=query) | Q(ip_address__icontains=query))
    if user_filter:
        logs = logs.filter(user_id=user_filter)

    total_count = logs.count()
    paginator = Paginator(logs, 15)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    users = User.objects.all()
    context = {
        'logs': page_obj,
        'page_obj': page_obj,
        'paginator': paginator,
        'total_count': total_count,
        'users': users,
        'query': query,
        'selected_user': user_filter,
    }
    return render(request, 'login_tracker/login_activity.html', context)


import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from django.http import HttpResponse
from django.utils import timezone


@manager_required
def export_login_excel(request):
    logs = LoginActivity.objects.select_related('user').all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Login Activity"

    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    headers = ['#', 'Username', 'Login Time', 'Logout Time', 'Duration', 'IP Address', 'User Agent']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 22
    for i, w in enumerate([5, 16, 22, 22, 12, 18, 40], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w

    for row_num, log in enumerate(logs, 2):
        row = [
            row_num - 1,
            log.user.username,
            log.login_time.strftime('%Y-%m-%d %H:%M') if log.login_time else '',
            log.logout_time.strftime('%Y-%m-%d %H:%M') if log.logout_time else 'Active',
            log.duration or 'Active',
            log.ip_address or '',
            (log.user_agent or '')[:80],
        ]
        for col, val in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.alignment = Alignment(horizontal='center' if col not in (2, 7) else 'left', vertical='center')

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="login_activity.xlsx"'
    wb.save(response)
    return response


@manager_required
def export_login_pdf(request):
    logs = LoginActivity.objects.select_related('user').all()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.75 * inch, bottomMargin=0.75 * inch)
    styles = getSampleStyleSheet()
    story = []

    title_s = ParagraphStyle('T', parent=styles['Title'], fontSize=18,
                              textColor=colors.HexColor('#1E40AF'), spaceAfter=6)
    sub_s = ParagraphStyle('S', parent=styles['Normal'], fontSize=10,
                            textColor=colors.grey, spaceAfter=20)
    story.append(Paragraph("Login Activity Report", title_s))
    story.append(Paragraph(f"Generated: {timezone.now().strftime('%B %d, %Y %H:%M')}", sub_s))

    data = [['#', 'Username', 'Login Time', 'Logout Time', 'Duration', 'IP Address']]
    for i, log in enumerate(logs, 1):
        data.append([
            str(i),
            log.user.username,
            log.login_time.strftime('%Y-%m-%d %H:%M') if log.login_time else '',
            log.logout_time.strftime('%Y-%m-%d %H:%M') if log.logout_time else 'Active',
            log.duration or 'Active',
            log.ip_address or '',
        ])

    t = Table(data, colWidths=[0.3*inch, 1.1*inch, 1.5*inch, 1.5*inch, 0.9*inch, 1.3*inch])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E40AF')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (1, 1), (1, -1), 'LEFT'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t)

    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="login_activity.pdf"'
    return response


@manager_required
def clear_login_activity(request):
    """Permanently delete all login activity records."""
    count = LoginActivity.objects.count()
    if request.method == 'POST':
        LoginActivity.objects.all().delete()
        from django.contrib import messages
        messages.success(request, f'{count} login record(s) have been permanently cleared.')
        return redirect('login_activity')
    return render(request, 'login_tracker/clear_login_confirm.html', {'count': count})


@manager_required
def live_search_login(request):
    """AJAX endpoint for login activity live search."""
    from django.http import JsonResponse
    from django.db.models import Q
    query   = request.GET.get('q', '').strip()
    user_id = request.GET.get('user', '').strip()

    logs = LoginActivity.objects.select_related('user').all()

    if query:
        logs = logs.filter(Q(user__username__icontains=query) | Q(ip_address__icontains=query))
    if user_id:
        logs = logs.filter(user_id=user_id)

    data = {
        'logs': [
            {
                'username':    log.user.username,
                'login_date':  log.login_time.strftime('%b %d, %Y') if log.login_time else '—',
                'login_time':  log.login_time.strftime('%I:%M %p') if log.login_time else '—',
                'logout_date': log.logout_time.strftime('%b %d, %Y') if log.logout_time else None,
                'logout_time': log.logout_time.strftime('%I:%M %p') if log.logout_time else None,
                'duration':    log.duration or None,
                'ip_address':  log.ip_address or '',
            }
            for log in logs[:100]
        ]
    }
    return JsonResponse(data)
