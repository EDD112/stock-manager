from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib import messages
from django.db.models import Q
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from activity.models import ActivityLog
from inventory.models import Product
from accounts.views import manager_required


@manager_required
def activity_log(request):
    logs = ActivityLog.objects.select_related('product', 'user').all()

    # Filters
    query = request.GET.get('q', '')
    user_id = request.GET.get('user', '')
    product_id = request.GET.get('product', '')
    action_filter = request.GET.get('action', '')

    if query:
        logs = logs.filter(Q(product_name__icontains=query) | Q(notes__icontains=query))
    if user_id:
        logs = logs.filter(user_id=user_id)
    if product_id:
        logs = logs.filter(product_id=product_id)
    if action_filter:
        logs = logs.filter(action=action_filter)

    total_count = logs.count()
    paginator = Paginator(logs, 15)
    page_number = request.GET.get('page', 1)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    users = User.objects.all()
    products = Product.objects.all()

    context = {
        'logs': page_obj,
        'page_obj': page_obj,
        'paginator': paginator,
        'total_count': total_count,
        'users': users,
        'products': products,
        'query': query,
        'selected_user': user_id,
        'selected_product': product_id,
        'selected_action': action_filter,
    }
    return render(request, 'activity/activity_log.html', context)


import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from django.http import HttpResponse
from django.utils import timezone


@manager_required
def export_activity_excel(request):
    logs = ActivityLog.objects.select_related('product', 'user').all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Activity Logs"

    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    headers = ['#', 'User', 'Product', 'Action', 'Quantity', 'Notes', 'Date & Time']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 22
    for i, w in enumerate([5, 16, 24, 12, 10, 36, 22], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w

    action_colors = {'ADD': 'DCFCE7', 'REMOVE': 'FEE2E2', 'UPDATE': 'DBEAFE', 'DELETE': 'FEF3C7'}
    for row_num, log in enumerate(logs, 2):
        row = [
            row_num - 1,
            log.user.username,
            log.get_product_name(),
            log.get_action_display(),
            log.quantity,
            log.notes or '',
            log.date.strftime('%Y-%m-%d %H:%M'),
        ]
        for col, val in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.alignment = Alignment(horizontal='center' if col not in (2, 3, 6) else 'left', vertical='center')
            if col == 4:
                cell.fill = PatternFill("solid", fgColor=action_colors.get(log.action, 'FFFFFF'))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="activity_logs.xlsx"'
    wb.save(response)
    return response


@manager_required
def export_activity_pdf(request):
    logs = ActivityLog.objects.select_related('product', 'user').all()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.75 * inch, bottomMargin=0.75 * inch)
    styles = getSampleStyleSheet()
    story = []

    title_s = ParagraphStyle('T', parent=styles['Title'], fontSize=18,
                              textColor=colors.HexColor('#1E40AF'), spaceAfter=6)
    sub_s = ParagraphStyle('S', parent=styles['Normal'], fontSize=10,
                            textColor=colors.grey, spaceAfter=20)
    story.append(Paragraph("Activity Log Report", title_s))
    story.append(Paragraph(f"Generated: {timezone.now().strftime('%B %d, %Y %H:%M')}", sub_s))

    data = [['#', 'User', 'Product', 'Action', 'Qty', 'Notes', 'Date']]
    for i, log in enumerate(logs, 1):
        data.append([
            str(i), log.user.username, log.get_product_name(),
            log.get_action_display(), str(log.quantity),
            (log.notes or '')[:40], log.date.strftime('%Y-%m-%d %H:%M'),
        ])

    t = Table(data, colWidths=[0.3*inch, 1.0*inch, 1.4*inch, 0.8*inch, 0.4*inch, 2.0*inch, 1.2*inch])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E40AF')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (2, 1), (2, -1), 'LEFT'),
        ('ALIGN', (5, 1), (5, -1), 'LEFT'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t)

    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="activity_logs.pdf"'
    return response


@manager_required
def clear_activity_logs(request):
    """Permanently delete all activity logs."""
    count = ActivityLog.objects.count()
    if request.method == 'POST':
        ActivityLog.objects.all().delete()
        messages.success(request, f'{count} activity log(s) have been permanently cleared.')
        return redirect('activity_log')
    return render(request, 'activity/clear_activity_confirm.html', {'count': count})


@manager_required
def live_search_activity(request):
    """AJAX endpoint for activity log live search."""
    from django.http import JsonResponse
    query      = request.GET.get('q', '').strip()
    user_id    = request.GET.get('user', '').strip()
    product_id = request.GET.get('product', '').strip()
    action_f   = request.GET.get('action', '').strip()

    logs = ActivityLog.objects.select_related('product', 'user').all()

    if query:
        logs = logs.filter(Q(product_name__icontains=query) | Q(notes__icontains=query))
    if user_id:
        logs = logs.filter(user_id=user_id)
    if product_id:
        logs = logs.filter(product_id=product_id)
    if action_f:
        logs = logs.filter(action=action_f)

    data = {
        'logs': [
            {
                'date':         log.date.strftime('%b %d, %Y'),
                'time':         log.date.strftime('%I:%M %p'),
                'username':     log.user.username,
                'product_name': log.product.name if log.product else (log.product_name or '(Deleted Product)'),
                'product_pk':   log.product.pk if log.product else None,
                'action':       log.action,
                'quantity':     log.quantity,
                'notes':        log.notes or '',
            }
            for log in logs[:100]
        ]
    }
    return JsonResponse(data)
